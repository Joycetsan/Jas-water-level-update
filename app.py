import streamlit as st
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill, Alignment, Border
from openpyxl.formula.translate import Translator
from datetime import datetime
import os
import warnings
import re
import io

# 忽略 openpyxl 条件格式警告
warnings.filterwarnings('ignore', category=UserWarning)

# 设置网页标题与布局
st.set_page_config(page_title="水位表自动化更新工具", layout="wide")

# =========================================================
# 工具函数
# =========================================================

def safe_write_cell(ws, cell_coord, value):
    for merged_range in ws.merged_cells.ranges:
        if cell_coord in merged_range:
            top_left = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            top_left.value = value
            return
    ws[cell_coord].value = value

def clone_stored_style(stored_style_dict, target_cell):
    if not stored_style_dict:
        return
    if stored_style_dict.get('font'):
        target_cell.font = Font.from_tree(stored_style_dict['font'].to_tree())
    if stored_style_dict.get('fill'):
        target_cell.fill = PatternFill.from_tree(stored_style_dict['fill'].to_tree())
    if stored_style_dict.get('border'):
        target_cell.border = Border.from_tree(stored_style_dict['border'].to_tree())
    if stored_style_dict.get('alignment'):
        target_cell.alignment = Alignment.from_tree(stored_style_dict['alignment'].to_tree())
    target_cell.number_format = stored_style_dict.get('number_format', 'General')

def load_fba_data(fba_file_bytes):
    """从上传的 FBA 字节流中读取数据"""
    fba_lookup = {}
    try:
        wb_fba = openpyxl.load_workbook(io.BytesIO(fba_file_bytes), data_only=True)
        if 'Sheet2' in wb_fba.sheetnames:
            ws_fba = wb_fba['Sheet2']
        else:
            ws_fba = wb_fba.active

        for row in ws_fba.iter_rows(min_row=4, max_col=2):
            asin = str(row[0].value).strip() if row[0].value else None
            qty = row[1].value if row[1].value is not None else 0
            if asin and len(asin) > 5:
                fba_lookup[asin] = qty
        return fba_lookup, None
    except Exception as e:
        return {}, str(e)

# =========================================================
# 核心处理逻辑
# =========================================================

def process_water_level_file(file_bytes, file_name, tasks, fba_lookup):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    wb_data = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    run_task1 = "任务1" in tasks
    run_task2 = "任务2" in tasks
    run_task3 = "任务3" in tasks

    moved_d_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    moved_d_font = Font(name='Bahnschrift', color="000000", size=8)
    mo_column_font = Font(name='等线', size=10)
    center_alignment = Alignment(horizontal='center', vertical='center')
    left_alignment = Alignment(horizontal='left', vertical='center')
    no_fill = PatternFill(fill_type=None)
    f_column_style = Font(name='Bahnschrift', bold=True, color="FF0000", size=8)
    e_column_style = Font(name='等线', bold=False, color="000000", size=10)

    today_dt = datetime.now()
    today_date = today_dt.date()
    exclude_list = ['合并', '成本', '数据源', 'ABA', '模板', 'ERP货件数据']
    target_columns = [3, 4] + list(range(10, 16))

    processed_count = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws_data_sheet = wb_data[sheet_name]

        if ws.sheet_state != 'visible' or any(kw in sheet_name for kw in exclude_list):
            continue

        today_row_idx = None
        for idx, row_cells in enumerate(ws_data_sheet.iter_rows(min_col=2, max_col=2), start=1):
            cell_val = row_cells[0].value
            if isinstance(cell_val, datetime) and cell_val.date() == today_date:
                today_row_idx = idx
                break

        if not today_row_idx:
            continue

        # 任务1：更新库存及E列公式
        if run_task1:
            target_asin = str(ws_data_sheet['J3'].value).strip() if ws_data_sheet['J3'].value else None
            if target_asin in fba_lookup:
                val_to_fill = fba_lookup[target_asin]
                target_f_cell = ws.cell(row=today_row_idx, column=6)
                target_f_cell.value = val_to_fill
                target_f_cell.font = f_column_style

                start_formula_row = today_row_idx - 8
                if start_formula_row > 0:
                    source_cell = ws.cell(row=start_formula_row, column=5)
                    source_formula = source_cell.value
                    if isinstance(source_formula, str) and source_formula.startswith('='):
                        origin_coord = source_cell.coordinate
                        for fill_idx in range(start_formula_row + 1, today_row_idx):
                            target_cell = ws.cell(row=fill_idx, column=5)
                            trans_formula = Translator(source_formula, origin=origin_coord).translate_formula(target_cell.coordinate)
                            target_cell.value = trans_formula
                            target_cell.font = e_column_style

            freeze_start_row = today_row_idx - 40
            freeze_end_row = today_row_idx - 14
            for freeze_r in range(freeze_start_row, freeze_end_row + 1):
                if freeze_r >= 2:
                    calculated_val = ws_data_sheet.cell(row=freeze_r, column=5).value
                    target_cell = ws.cell(row=freeze_r, column=5)
                    if not isinstance(target_cell, MergedCell):
                        target_cell.value = calculated_val

        # J列向下填充
        last_row = ws_data_sheet.max_row
        row_cursor = max(2, today_row_idx - 30)
        while row_cursor <= last_row:
            continuous_rows = []
            check_r = row_cursor
            while check_r <= last_row and ws_data_sheet.cell(row=check_r, column=11).value is not None:
                continuous_rows.append(check_r)
                check_r += 1

            if len(continuous_rows) >= 3:
                first_j_val = ws.cell(row=continuous_rows[0], column=10).value
                if first_j_val:
                    for fill_r in continuous_rows[1:]:
                        if ws.cell(row=fill_r, column=10).value in [None, ""]:
                            ws.cell(row=fill_r, column=10).value = first_j_val
                row_cursor = check_r
            else:
                row_cursor += 1

        # 任务2：扫描与填充
        if run_task2:
            for r in range(today_row_idx + 1, last_row + 1):
                n_raw = ws_data_sheet.cell(row=r, column=14).value
                if isinstance(n_raw, str) and n_raw.startswith("#"): continue
                try:
                    n_val = float(n_raw) if n_raw is not None else 0.0
                    if n_val != 0:
                        c_val_orig = ws_data_sheet.cell(row=r, column=3).value
                        if c_val_orig in [None, ""]:
                            d_cell = ws.cell(row=r, column=4)
                            if not isinstance(d_cell, MergedCell):
                                d_cell.value = n_val
                                d_cell.fill = moved_d_fill
                                d_cell.font = moved_d_font
                                d_cell.alignment = center_alignment
                                ws_data_sheet.cell(row=r, column=4).value = n_val
                except (ValueError, TypeError):
                    continue

            tasks_move_back, tasks_move_forward, tasks_freeze_history = [], [], []

            start_r = max(2, today_row_idx - 20)
            for r in range(start_r, today_row_idx):
                n_raw = ws_data_sheet.cell(row=r, column=14).value
                o_raw = ws_data_sheet.cell(row=r, column=15).value
                if isinstance(n_raw, str) and n_raw.startswith("#"): continue
                if isinstance(o_raw, str) and o_raw.startswith("#"): continue
                try:
                    n_v = float(n_raw) if n_raw is not None else 0.0
                    o_v = float(o_raw) if o_raw is not None else 0.0
                    if o_v >= n_v and n_v > 0:
                        freeze_cols = [10, 11, 13, 14, 15]
                        freeze_data = {col: ws_data_sheet.cell(row=r, column=col).value for col in freeze_cols}
                        tasks_freeze_history.append({'row': r, 'data': freeze_data})
                except:
                    continue	

            for r in range(2, today_row_idx + 1):
                n_raw = ws_data_sheet.cell(row=r, column=14).value
                o_raw = ws_data_sheet.cell(row=r, column=15).value
                if isinstance(n_raw, str) and n_raw.startswith("#"): continue
                if isinstance(o_raw, str) and o_raw.startswith("#"): continue
                try:
                    n_val = float(n_raw) if n_raw is not None else 0.0
                    o_val = float(o_raw) if o_raw is not None else 0.0
                    if n_val > 0 and o_val <= (0.3 * n_val):
                        row_data, row_styles = {}, {}
                        for col in target_columns:
                            src_cell = ws.cell(row=r, column=col)
                            row_data[col] = src_cell.value
                            if src_cell.has_style:
                                row_styles[col] = {
                                    'font': src_cell.font, 'fill': src_cell.fill,
                                    'border': src_cell.border, 'alignment': src_cell.alignment,
                                    'number_format': src_cell.number_format
                                }
                        tasks_move_back.append({'old_row': r, 'data': row_data, 'styles': row_styles})
                except:
                    continue

            for r in range(today_row_idx, last_row + 1):
                o_raw = ws_data_sheet.cell(row=r, column=15).value
                n_raw = ws_data_sheet.cell(row=r, column=14).value
                if isinstance(n_raw, str) and n_raw.startswith("#"): continue
                if isinstance(o_raw, str) and o_raw.startswith("#"): continue
                try:
                    o_val = float(o_raw) if o_raw is not None else 0.0
                    n_val = float(n_raw) if n_raw is not None else 0.0
                    if n_val > 0 and o_val > (2.0 / 3.0 * n_val):
                        row_data, row_styles = {}, {}
                        for col in target_columns:
                            src_cell = ws.cell(row=r, column=col)
                            row_data[col] = ws_data_sheet.cell(row=r, column=col).value if col in [13, 14, 15] else src_cell.value
                            if src_cell.has_style:
                                row_styles[col] = {
                                    'font': src_cell.font, 'fill': src_cell.fill,
                                    'border': src_cell.border, 'alignment': src_cell.alignment,
                                    'number_format': src_cell.number_format
                                }
                        tasks_move_forward.append({'old_row': r, 'data': row_data, 'styles': row_styles})
                except (ValueError, TypeError):
                    continue

            affected_rows = [t['old_row'] for t in tasks_move_back] + [t['old_row'] for t in tasks_move_forward]
            for r in affected_rows:
                for col in target_columns:
                    cell = ws.cell(row=r, column=col)
                    if not isinstance(cell, MergedCell):
                        cell.value = None
                        cell.fill = no_fill

            for task in tasks_freeze_history:
                for col, val in task['data'].items():
                    target_cell = ws.cell(row=task['row'], column=col)
                    if not isinstance(target_cell, MergedCell):
                        target_cell.value = val

            def find_next_empty_row(ws, start_row, direction, check_col=10):
                r = start_row
                max_row = ws.max_row
                while True:
                    if r < 2: return None
                    if r > max_row: return max_row + 1
                    cell_val = ws.cell(row=r, column=check_col).value
                    if cell_val in [None, ""]: return r
                    r += direction

            for i, task in enumerate(tasks_move_back):
                target_start_row = today_row_idx + 3 + i
                new_r = find_next_empty_row(ws, target_start_row, direction=1)
                if new_r is None: continue
                for col, val in task['data'].items():
                    target_cell = ws.cell(row=new_r, column=col)
                    if isinstance(target_cell, MergedCell): continue
                    if isinstance(val, str) and val.startswith("="):
                        origin_coord = ws.cell(row=task['old_row'], column=col).coordinate
                        target_cell.value = Translator(val, origin=origin_coord).translate_formula(target_cell.coordinate)
                    else:
                        target_cell.value = val
                    if col in [3, 4]:
                        clone_stored_style(task['styles'].get(col), target_cell)
                    else:
                        if col in [13, 14, 15]:
                            target_cell.font = mo_column_font
                            target_cell.alignment = left_alignment if col == 13 else center_alignment

            for i, task in enumerate(tasks_move_forward):
                target_start_row = today_row_idx - 1 - i
                new_r = find_next_empty_row(ws, target_start_row, direction=-1)
                if new_r is None: continue
                for col, val in task['data'].items():
                    target_cell = ws.cell(row=new_r, column=col)
                    if isinstance(target_cell, MergedCell): continue
                    if col in [13, 14, 15]:
                        target_cell.value = ws_data_sheet.cell(row=task['old_row'], column=col).value
                    else:
                        target_cell.value = val
                    if col in [3, 4]:
                        clone_stored_style(task['styles'].get(col), target_cell)
                    else:
                        if col in [13, 14, 15]:
                            target_cell.font = mo_column_font
                            target_cell.alignment = left_alignment if col == 13 else center_alignment

        # 任务3：汇总统计
        if run_task3:
            sum_with_bg = 0
            sum_colored_font = 0
            stat_columns = [3, 4]

            for row in range(4, ws.max_row + 1):
                date_val = ws_data_sheet.cell(row=row, column=2).value
                if isinstance(date_val, datetime) and date_val.date() >= today_date:
                    for col in stat_columns:
                        cell_style = ws.cell(row=row, column=col)
                        val = cell_style.value
                        if isinstance(val, (int, float)):
                            has_bg = False
                            if cell_style.fill and cell_style.fill.fill_type not in [None, 'none']:
                                color_obj = cell_style.fill.start_color
                                if color_obj.index != '00000000' and color_obj.rgb != '00000000':
                                    has_bg = True

                            is_not_black = False
                            if not has_bg:
                                font_color = cell_style.font.color
                                if font_color:
                                    rgb = str(font_color.rgb)
                                    if rgb not in ['00000000', 'FF000000', 'None'] and font_color.type != 'theme':
                                        is_not_black = True
                                    elif font_color.type == 'theme' and font_color.theme != 1:
                                        is_not_black = True

                            if has_bg:
                                sum_with_bg += val
                            elif is_not_black:
                                sum_colored_font += val

            safe_write_cell(ws, 'E3', sum_with_bg)
            safe_write_cell(ws, 'F3', sum_colored_font)

        processed_count += 1

    # 将结果保存到内存流并返回
    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream.getvalue(), processed_count

# =========================================================
# Streamlit UI 页面渲染
# =========================================================

st.title("🚀 水位表自动化更新工具")
st.markdown("请在左侧边栏输入密钥解锁系统功能。")

# ---------------------------------------------------------
# 1. 侧边栏配置：仅保留密码验证输入框
# ---------------------------------------------------------
with st.sidebar:
    st.header("⚙️ 系统配置")
    access_code = st.text_input("请输入访问密钥", type="password")

# ---------------------------------------------------------
# 2. 校验密钥（密码正确才继续往下渲染，否则终止）
# ---------------------------------------------------------
if access_code != "ltmx2026":
    st.warning("🔒 请在左侧边栏输入正确的访问密钥以解锁功能。")
    
    # 底部页脚（未解锁前也展示在最下方）
    st.markdown("---")
    st.markdown(
        "<div style='text-align: center; color: #888888; font-size: 14px; padding: 10px;'>"
        "Copyright © Jas | Contact: <a href='mailto:Joyce.qinyizhang@Gmail.com' style='color: #888888;'>Joyce.qinyizhang@Gmail.com</a>"
        "</div>",
        unsafe_allow_html=True
    )
    st.stop()

# ---------------------------------------------------------
# 3. 密码正确后：在侧边栏显示任务选择框
# ---------------------------------------------------------
with st.sidebar:
    st.divider()
    
    # 任务选择框
    task_options_map = {
        "任务1": "任务1",
        "任务2": "任务2",
        "任务3": "任务3"
    }
    
    selected_task_labels = st.multiselect(
        "选择要执行的任务",
        options=list(task_options_map.keys()),
        default=list(task_options_map.keys()),
        help="可以根据需要勾选单个或多个任务"
    )
    
    selected_tasks = [task_options_map[label] for label in selected_task_labels]

# ---------------------------------------------------------
# 4. 密码正确后：在主界面显示文档链接与任务规则说明
# ---------------------------------------------------------
# 新增：说明文档链接提示（可以将下面的 xxx 替换为你的文档实际链接，如飞书/钉钉/Notion/OneDrive链接）
st.info("📄 **详见说明文档**：[点击此处查看完整 SOP 使用指南与文档](https://jcn3bijgp23x.feishu.cn/wiki/KrwUwFGLyi9HSekuE8DcCYiynpg)")

with st.expander("📖 点击查看【任务 1、2、3】具体处理规则说明", expanded=False):
    st.markdown("""
    * **[1] 任务 1：更新今日库存及销售列公式填充**
      * 匹配每个sheet J3 的 ASIN 的对应 **FBA** 库存并更新至今天行
      * 自动向下填充今天行之前的 E 列即 **销售列** 的公式
      * 对历史今天前 **第 14-40 行** 的销售列进行值数据封存（防止后续变动影响历史数据）
    
    ---
    
    * **[2] 任务 2：FBA货件移动**
      * 对 **已接收的货件** 的数据封存到今天之前
      *  **未接收的货件** 移到今天后的两行之后
      * 已发货且贴了FBA号的货件在到仓列格式改为在途的格式
    
    ---
    
    * **[3] 任务 3：汇总统计（颜色分类统计写入 E3, F3）**
      * 今日及以后的排的 **在途和采购** 数据会统计汇总填入每个sheet对应的 E3 和 F3 单元格
    """)

# ---------------------------------------------------------
# 5. 密码正确后：在主界面显示文件上传与处理按钮
# ---------------------------------------------------------
col1, col2 = st.columns(2)

with col1:
    fba_file = st.file_uploader("1. 上传【FBA仓库明细.xlsx】(任务1必备)", type=["xlsx"])

with col2:
    water_files = st.file_uploader("2. 上传【水位表文件】(可多选)", type=["xlsx"], accept_multiple_files=True)

# 处理与提交
if st.button("▶️ 开始在线处理", type="primary"):
    if not water_files:
        st.error("❌ 请至少上传一个水位表文件！")
        st.stop()

    fba_lookup = {}
    if "任务1" in selected_tasks:
        if not fba_file:
            st.error("❌ 已勾选 [任务1]，必须上传 FBA仓库明细.xlsx 文件！")
            st.stop()
        
        fba_lookup, err = load_fba_data(fba_file.read())
        if err:
            st.error(f"❌ 读取 FBA 文件失败: {err}")
            st.stop()
        else:
            st.success(f"📊 成功读取 FBA 明细，包含 {len(fba_lookup)} 条 ASIN 库存数据")

    # 开始循环处理水位表
    today_str = datetime.now().strftime('%Y%m%d')
    
    st.divider()
    st.subheader("📦 处理结果下载")

    for file_obj in water_files:
        file_bytes = file_obj.read()
        file_name = file_obj.name

        with st.spinner(f"正在处理：{file_name} ..."):
            try:
                processed_bytes, count = process_water_level_file(
                    file_bytes, file_name, selected_tasks, fba_lookup
                )
                
                # 清洗文件名格式
                base_name = os.path.splitext(file_name)[0]
                base_name = re.sub(r'_\d{8}$', '', base_name)
                download_filename = f"{base_name}_{today_str}.xlsx"

                st.success(f"✅ 【{file_name}】处理成功（处理页签数：{count}）")
                st.download_button(
                    label=f"⬇️ 下载处理后的 {download_filename}",
                    data=processed_bytes,
                    file_name=download_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=download_filename
                )
            except Exception as e:
                st.error(f"❌ 文件 【{file_name}】 处理失败，错误原因为: {str(e)}")

# ---------------------------------------------------------
# 6. 解锁后的页面最下方版权与联系信息 (Footer)
# ---------------------------------------------------------
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: #888888; font-size: 14px; padding: 10px;'>"
    "Copyright © Jas"
    "</div>",
    unsafe_allow_html=True
)